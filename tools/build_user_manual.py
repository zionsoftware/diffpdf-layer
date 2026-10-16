"""画面付きMarkdown原稿から、画像内包の配布用DOCXを生成・検査する。

依存はtools/requirements-docs.txt。比較ツールからは呼び出さない。
扱うMarkdown: 見出し、段落、表、番号/箇条書き、コード、強調、リンク、PNG。
未対応のブロック構文は黙って落とさずエラーにする。
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import struct
from urllib.parse import urlsplit
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "Meiryo"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
      "r": "http://schemas.openxmlformats.org/package/2006/relationships"}
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)\Z")


@dataclass
class Block:
    kind: str
    value: object


def source_digest(source):
    """Gitの改行変換だけで原稿の変更と判定しないようLFで照合する。"""
    return hashlib.sha256(source.read_text(encoding="utf-8").encode("utf-8")).hexdigest()


def table_cells(line):
    return [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def parse_markdown(text):
    """この原稿で使用するブロックを明示的に解析する。HTMLは許容しない。"""
    lines, blocks, i = text.splitlines(), [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            language, code = line[3:].strip(), []
            i += 1
            while i < len(lines) and lines[i] != "```":
                code.append(lines[i]); i += 1
            if i == len(lines):
                raise ValueError("コードブロックが閉じていません")
            blocks.append(Block("code", (language, "\n".join(code))))
            i += 1
            continue
        heading = re.fullmatch(r"(#{1,6}) (.+)", line)
        picture = IMAGE.fullmatch(line)
        item = re.fullmatch(r"(\d+\. |[-*] )(.+)", line)
        if heading:
            blocks.append(Block("heading", (len(heading[1]), heading[2])))
        elif picture:
            blocks.append(Block("image", (picture[1], picture[2])))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(table_cells(lines[i])); i += 1
            if len(rows) < 2 or not all(re.fullmatch(r":?-+:?", c) for c in rows[1]):
                raise ValueError("表の区切り行が不正です")
            rows.pop(1)
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError("表の列数が一致しません")
            blocks.append(Block("table", rows))
            continue
        elif item:
            ordered, items = item[1][0].isdigit(), []
            while i < len(lines):
                match = re.fullmatch(r"(\d+\. |[-*] )(.+)", lines[i])
                if not match or match[1][0].isdigit() != ordered:
                    break
                items.append(match[2]); i += 1
            blocks.append(Block("list", (ordered, items)))
            continue
        else:
            paragraph = []
            while i < len(lines) and lines[i].strip():
                if re.match(r"^(?:#|>|<|!\[|\||```|[-*] |\d+\. )", lines[i]):
                    if not paragraph:
                        raise ValueError(f"未対応のMarkdown構文: {lines[i]}")
                    break
                paragraph.append(lines[i]); i += 1
            blocks.append(Block("paragraph", " ".join(paragraph)))
            continue
        i += 1
    return blocks


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return re.sub(r"\*\*(.+?)\*\*|`([^`]+)`|\*([^*]+)\*", lambda m: next(x for x in m.groups() if x is not None), text)


def element(tag, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn("w:" + key), str(value))
    return el


def font_style(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = element("w:rFonts"); props.append(fonts)
    for key in list(fonts.attrib):
        if "theme" in key.lower():
            del fonts.attrib[key]
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), FONT)
    lang = element("w:lang", val="ja-JP", eastAsia="ja-JP")
    props.append(lang)
    for name in ("szCs", "iCs"):
        old = props.find(qn("w:"+name))
        if old is not None: props.remove(old)
    props.append(element("w:szCs",val=int(size*2)))
    props.append(element("w:iCs",val="false"))
    pprops=style.element.get_or_add_pPr()
    for child in list(pprops):
        if child.tag==qn("w:pBdr") or style.name=="Subtitle" and child.tag==qn("w:numPr"):
            pprops.remove(child)


def hyperlink(paragraph, label, target, *, anchor=False):
    link = OxmlElement("w:hyperlink")
    if anchor:
        link.set(qn("w:anchor"), target)
    else:
        rid = paragraph.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    props.append(element("w:color", val="1F4E79"))
    props.append(element("w:u", val="single"))
    run.append(props)
    value = OxmlElement("w:t"); value.text = label; run.append(value)
    link.append(run); paragraph._p.append(link)


def inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            label, target = link.groups()
            if not urlsplit(target).scheme:
                # 配布先に他のMDがなくても参照先を特定できる補足リンク。
                target = "https://github.com/zionsoftware/diffpdf-layer/blob/main/docs/" + target
            hyperlink(paragraph, label, target)
        else:
            run = paragraph.add_run(plain(part))
            if part.startswith("**"):
                run.bold = True
            elif part.startswith("*"):
                run.italic = True
            elif part.startswith("`"):
                run.font.name = "Consolas"


def png(source, target):
    path = (source.parent / target).resolve()
    if urlsplit(target).scheme or not path.is_relative_to(source.parent.resolve()):
        raise ValueError("画像は原稿のフォルダー以下のローカルPNGに限定します")
    data = path.read_bytes()
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"PNGではありません: {path}")
    return path, data, struct.unpack(">II", data[16:24])


def numbering(doc, ordered):
    root = doc.part.numbering_part.element
    num_ids = [int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))]
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in root.findall(qn("w:abstractNum"))]
    num_id, abstract_id = max(num_ids, default=0)+1, max(abstract_ids, default=0)+1
    abstract = element("w:abstractNum", abstractNumId=abstract_id)
    level = element("w:lvl", ilvl=0)
    for tag, attrs in [("w:start", {"val":1}), ("w:numFmt", {"val":"decimal" if ordered else "bullet"}), ("w:lvlText", {"val":"%1." if ordered else "•"})]:
        level.append(element(tag, **attrs))
    props = OxmlElement("w:pPr"); props.append(element("w:ind", left=360, hanging=260)); level.append(props)
    abstract.append(level); root.append(abstract)
    num = element("w:num", numId=num_id); num.append(element("w:abstractNumId", val=abstract_id)); root.append(num)
    return num_id


def make_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = False
    widths = [52, 114] if len(rows[0]) == 2 else [166/len(rows[0])] * len(rows[0])
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element("w:"+edge, val="single", sz=4, color="D9D9D9"))
    props.append(borders)
    for index, values in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        tr = row._tr.get_or_add_trPr(); tr.append(element("w:cantSplit"))
        if index == 0:
            tr.append(element("w:tblHeader"))
        for cell, width, text in zip(row.cells, widths, values):
            cell.width = Mm(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margins.append(element("w:"+side, w=80, type="dxa"))
            cell._tc.get_or_add_tcPr().append(margins)
            if index == 0:
                cell._tc.get_or_add_tcPr().append(element("w:shd", fill="E8EFF6"))
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            inline(p, text)
            for run in p.runs:
                run.font.size = Pt(10)
                if index == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def body_blocks(blocks):
    return [b for i,b in enumerate(blocks) if i != 0 and not (b.kind == "paragraph" and re.fullmatch(r"対象: diffpdf-layer .+。更新日: .+。",str(b.value)))]


def build(source, output):
    text = source.read_text(encoding="utf-8")
    blocks = parse_markdown(text)
    if not blocks or blocks[0] != Block("heading", (1,"ユーザーマニュアル")):
        raise ValueError("先頭の見出しはユーザーマニュアルとしてください")
    meta = re.search(r"対象: diffpdf-layer ([^。]+)。更新日: ([0-9-]+)。", text)
    if not meta: raise ValueError("原稿の対象バージョン・更新日がありません")
    version, date = meta.groups()
    doc = Document(); section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    section.header_distance = section.footer_distance = Mm(10)
    section.different_first_page_header_footer = False
    for name, size, bold in [("Normal",10.5,False),("Title",24,False),("Subtitle",18,False),("Heading 1",16,True),("Heading 2",12,True),("Heading 3",11,True),("Caption",9,False),("List Number",10.5,False),("List Bullet",10.5,False)]:
        font_style(doc.styles[name],size,bold)
    normal=doc.styles["Normal"].paragraph_format
    normal.line_spacing=1.25; normal.space_after=Pt(5); normal.widow_control=True
    for name in ("Heading 1","Heading 2","Heading 3"):
        p=doc.styles[name].paragraph_format
        p.keep_with_next=True; p.keep_together=True; p.space_before=Pt(12); p.space_after=Pt(6)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("ページ ")
    field=element("w:fldSimple",instr="PAGE")
    run=OxmlElement("w:r"); t=OxmlElement("w:t");t.text="1";run.append(t);field.append(run);footer._p.append(field)
    doc.settings.element.append(element("w:updateFields",val="false"))
    cover=doc.add_paragraph("PDFレイヤー同一性検証ツール", "Title")
    cover.paragraph_format.space_before=Mm(50)
    doc.add_paragraph("ユーザーマニュアル", "Subtitle")
    doc.add_paragraph("diffpdf-layer")
    doc.add_paragraph(f"対象ソフトウェアバージョン  {version}")
    doc.add_paragraph(f"文書の更新日  {date}")
    doc.add_page_break()
    doc.add_paragraph("目次", "Title")
    doc.add_paragraph("各項目は章へのリンクです。WordではCtrlキーを押しながらクリックして移動できます。目次のフィールド更新は不要です。")
    chapters=[b.value[1] for b in blocks if b.kind=="heading" and b.value[0]==2]
    for index, heading in enumerate(chapters,1):
        p=doc.add_paragraph(); hyperlink(p,heading,f"manual_ch_{index}",anchor=True)
        p.paragraph_format.space_after=Pt(7)
    doc.add_page_break()
    body=doc._element.body
    body.insert(len(body)-1,element("w:bookmarkStart",id=900,name="manual_body_start"))
    chapter=0
    for block in body_blocks(blocks):
        if block.kind=="heading":
            level, title=block.value
            p=doc.add_paragraph(style=f"Heading {min(level-1,3)}");inline(p,title)
            if level==2:
                chapter+=1
                p._p.insert(0,element("w:bookmarkStart",id=chapter,name=f"manual_ch_{chapter}"))
                p._p.append(element("w:bookmarkEnd",id=chapter))
        elif block.kind=="paragraph":
            p=doc.add_paragraph();inline(p,block.value)
            if re.match(r"図\d+ ",block.value):
                p.style="Caption";p.paragraph_format.keep_together=True
        elif block.kind=="code":
            p=doc.add_paragraph();p.paragraph_format.keep_together=True
            run=p.add_run(block.value[1]);run.font.name="Consolas";run.font.size=Pt(10)
        elif block.kind=="list":
            ordered,items=block.value;num_id=numbering(doc,ordered)
            for item in items:
                p=doc.add_paragraph(style="List Number" if ordered else "List Bullet")
                n=OxmlElement("w:numPr");n.append(element("w:ilvl",val=0));n.append(element("w:numId",val=num_id));p._p.get_or_add_pPr().append(n)
                inline(p,item)
        elif block.kind=="table":
            make_table(doc,block.value)
        elif block.kind=="image":
            alt,target=block.value;path,data,(width,height)=png(source,target)
            physical=min(145 if width<=600 else 166,110*width/height)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
            picture=p.add_run().add_picture(str(path),width=Mm(physical))
            picture._inline.docPr.set("descr",alt)
    body.insert(len(body)-1,element("w:bookmarkEnd",id=900))
    props=doc.core_properties
    props.title="PDFレイヤー同一性検証ツール ユーザーマニュアル"
    props.subject="diffpdf-layer 操作マニュアル"
    props.author="";props.last_modified_by="";props.keywords="diffpdf-layer"
    props.created=props.modified=datetime.fromisoformat(date).replace(tzinfo=timezone.utc)
    props.identifier="sha256:"+source_digest(source)
    props.comments="Markdown原稿から生成。画像は文書内に埋め込み。"
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    return check(source,output)


def check(source,output):
    blocks=parse_markdown(source.read_text(encoding="utf-8"))
    expected=[]; images=[]
    for b in body_blocks(blocks):
        if b.kind=="heading": expected.append(plain(b.value[1]))
        elif b.kind=="paragraph": expected.append(plain(b.value))
        elif b.kind=="code": expected.append(b.value[1])
        elif b.kind=="list": expected.extend(plain(i) for i in b.value[1])
        elif b.kind=="table": expected.extend(plain(c) for row in b.value for c in row)
        elif b.kind=="image": images.append(png(source,b.value[1]))
    with ZipFile(output) as z:
        xml=ET.fromstring(z.read("word/document.xml"))
        actual=[];active=False
        for child in xml.find("w:body",NS):
            if child.tag==qn("w:bookmarkStart") and child.get(qn("w:name"))=="manual_body_start": active=True;continue
            if child.tag==qn("w:bookmarkEnd") and child.get(qn("w:id"))=="900": break
            if active:
                for p in ([child] if child.tag==qn("w:p") else child.findall(".//w:p",NS)):
                    value="".join((t.text or "") if t.tag==qn("w:t") else "\n" if t.tag==qn("w:br") else "\t"
                                  for t in p.iter() if t.tag in {qn("w:t"),qn("w:br"),qn("w:tab")})
                    if value: actual.append(value)
        if actual!=expected:
            raise ValueError("MarkdownとWordの本文が一致しません")
        rels=ET.fromstring(z.read("word/_rels/document.xml.rels"))
        image_rels=[r for r in rels if r.get("Type","").endswith("/image")]
        if len(image_rels)!=len(images) or any(r.get("TargetMode")=="External" for r in image_rels):
            raise ValueError("画像の数または埋め込み方法が不正です")
        embedded={hashlib.sha256(z.read("word/"+r.get("Target"))).hexdigest() for r in image_rels}
        if embedded!={hashlib.sha256(data).hexdigest() for _,data,_ in images}:
            raise ValueError("埋め込み画像がPNG原本と一致しません")
        sizes=xml.findall(".//wp:extent",NS)
        if len(sizes)!=len(images): raise ValueError("画像の表示数が不正です")
        for extent,(_,_,(w,h)) in zip(sizes,images):
            if abs(int(extent.get("cx"))/int(extent.get("cy"))-w/h)>0.0001:
                raise ValueError("画像の縦横比が不正です")
        page=xml.find(".//w:sectPr/w:pgSz",NS)
        if abs(int(page.get(qn("w:w")))-11906)>2 or abs(int(page.get(qn("w:h")))-16838)>2:
            raise ValueError("A4縦ではありません")
        headings=xml.findall('.//w:pPr/w:pStyle[@w:val="Heading1"]',NS)
        chapter_count=sum(b.kind=="heading" and b.value[0]==2 for b in blocks)
        if len(headings)!=chapter_count: raise ValueError("章の見出し数が不正です")
        anchors={a.get(qn("w:name")) for a in xml.findall(".//w:bookmarkStart",NS)}
        toc=[a.get(qn("w:anchor")) for a in xml.findall(".//w:hyperlink",NS) if a.get(qn("w:anchor"))]
        if len(toc)!=chapter_count or not set(toc)<=anchors: raise ValueError("目次リンクが不正です")
        section=xml.find(".//w:sectPr",NS)
        if section.find("w:titlePg",NS) is not None:
            raise ValueError("表紙にもページ番号を表示してください")
        footer_ref=section.find('w:footerReference[@w:type="default"]',NS)
        footer_target=next((r.get("Target") for r in rels if footer_ref is not None and r.get("Id")==footer_ref.get(qn("r:id"))),None)
        if not footer_target:
            raise ValueError("全ページ共通のフッターがありません")
        footer=ET.fromstring(z.read("word/"+footer_target))
        page_field=footer.find('.//w:fldSimple[@w:instr="PAGE"]',NS)
        alignment=footer.find('.//w:pPr/w:jc',NS)
        if page_field is None or alignment is None or alignment.get(qn("w:val"))!="center" or "ページ " not in "".join(footer.itertext()):
            raise ValueError("下中央のページ番号がありません")
        core=ET.fromstring(z.read("docProps/core.xml"))
        identifier=core.find("{http://purl.org/dc/elements/1.1/}identifier")
        if identifier is None or identifier.text!="sha256:"+source_digest(source):
            raise ValueError("原稿の更新がDOCXに反映されていません。再生成してください")
    return {"source":str(source),"output":str(output),"body_blocks_checked":len(expected),"chapters":chapter_count,"toc_links":len(toc),"embedded_images":len(images),"external_images":0,"a4_portrait":True,"source_sha256":source_digest(source)}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source",type=Path,default=ROOT/"docs/user-manual.md")
    parser.add_argument("--output",type=Path,default=ROOT/"docs/user-manual.docx")
    parser.add_argument("--check",action="store_true",help="既存DOCXを原稿・PNGと照合する（生成しない）")
    args=parser.parse_args()
    result=(check if args.check else build)(args.source.resolve(),args.output.resolve())
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
